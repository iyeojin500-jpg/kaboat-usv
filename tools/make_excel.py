"""결과 CSV → 엑셀 정리 파일 (results/결과정리.xlsx).

  python tools/make_excel.py                  # results/ 의 CSV 로 생성
  python tools/make_excel.py --results results_seed42 --out 정리.xlsx

시트
  결과정리      요청 표(구분·필요한 데이터·단위) 그대로 CASE1/CASE2 값과 개선율 — 모든 값은 아래 시트를 참조하는 수식
  시스템요약    system_summary.csv (항목 × CASE)
  AMR별         case1_amr.csv + case2_amr.csv
  주문별_CASE1  case1_orders.csv
  주문별_CASE2  case2_orders.csv
"""
from __future__ import annotations

import argparse
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="33415C")
GRP_FILL = PatternFill("solid", fgColor="EEF1F6")
SUB_FILL = PatternFill("solid", fgColor="F7F7F4")
THIN = Side(style="thin", color="C9CDD4")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def num(v: str):
    try:
        x = float(v)
        return int(x) if x.is_integer() and "." not in v else x
    except (TypeError, ValueError):
        return v


def read_csv(path):
    with open(path, encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    return rows[0], [[num(v) for v in r] for r in rows[1:]]


def data_sheet(wb, name, header, rows, widths=None):
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font, c.fill, c.alignment = f(True, "FFFFFF"), HDR_FILL, Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = f()
            if isinstance(c.value, float):
                c.number_format = "#,##0.0"
            elif isinstance(c.value, int):
                c.number_format = "#,##0"
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 32
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, min(26, len(str(h)) * 1.6)))
    return ws


def build(results: str, out: str):
    sys_h, sys_rows = read_csv(os.path.join(results, "system_summary.csv"))
    cases = [r[0] for r in sys_rows]
    assert cases[:2] == ["CASE1", "CASE2"], "system_summary.csv 에 CASE1, CASE2 가 있어야 합니다"
    ord_h, ord1 = read_csv(os.path.join(results, "case1_orders.csv"))
    _, ord2 = read_csv(os.path.join(results, "case2_orders.csv"))
    amr_h, amr1 = read_csv(os.path.join(results, "case1_amr.csv"))
    _, amr2 = read_csv(os.path.join(results, "case2_amr.csv"))

    wb = Workbook()
    main = wb.active
    main.title = "결과정리"
    # ---- 데이터 시트
    sys_t = [[k] + [r[i] for r in sys_rows] for i, k in enumerate(sys_h)]
    data_sheet(wb, "시스템요약", sys_t[0], sys_t[1:], {"CASE": 40})
    wb["시스템요약"].column_dimensions["A"].width = 40
    for c in "BCDE":
        wb["시스템요약"].column_dimensions[c].width = 26
    data_sheet(wb, "AMR별", ["CASE"] + amr_h, [["CASE1"] + r for r in amr1] + [["CASE2"] + r for r in amr2])
    data_sheet(wb, "주문별_CASE1", ord_h, ord1)
    data_sheet(wb, "주문별_CASE2", ord_h, ord2)

    # 데이터가 있는 범위만 참조 (전체 열 참조보다 빠르고 명확)
    n_ord = {"CASE1": len(ord1) + 1, "CASE2": len(ord2) + 1}                  # 마지막 행 번호
    n_amr = len(amr1) + len(amr2) + 1
    n_sys = len(sys_t)
    _oc = {h: get_column_letter(i + 1) for i, h in enumerate(ord_h)}

    def ORD(case):
        return "주문별_CASE1" if case == "CASE1" else "주문별_CASE2"

    def ocol(key, case):
        L = _oc[key]
        return f"{ORD(case)}!${L}$2:${L}${n_ord[case]}"

    ac = {h: get_column_letter(i + 2) for i, h in enumerate(amr_h)}            # AMR별 열 문자 (A=CASE)

    def SYS(key, case):
        col = "B" if case == "CASE1" else "C"
        return f'INDEX(시스템요약!${col}$2:${col}${n_sys},MATCH("{key}",시스템요약!$A$2:$A${n_sys},0))'

    def amr_sum(field, case, amr=None):
        col = ac[field]
        rng = f"AMR별!${col}$2:${col}${n_amr}"
        if amr:
            return f'SUMIFS({rng},AMR별!$A$2:$A${n_amr},"{case}",AMR별!$B$2:$B${n_amr},"{amr}")'
        return f'SUMIFS({rng},AMR별!$A$2:$A${n_amr},"{case}")'

    # 행 정의: (구분, 필요한 데이터, 단위, CASE→수식(=제외), 방향(-1 낮을수록 좋음 / +1 높을수록 / 0 비교 안 함), 출처, 비고, 표시형식)
    R = []
    T = "#,##0.0"
    N = "#,##0"
    rng = lambda key, case: (f'TEXT(MIN({ocol(key, case)}),"#,##0.0")&" ~ "&'
                             f'TEXT(MAX({ocol(key, case)}),"#,##0.0")')
    R += [("주문", "주문 발생시간", "초", lambda c: rng("발생시각(s)", c), 0, "주문별_CASE* · 발생시각(s)",
           "첫 주문 ~ 마지막 주문 (주문별 값은 주문별 시트)", "@"),
          ("주문", "주문 완료시간", "초", lambda c: rng("출고완료(s)", c), 0, "주문별_CASE* · 출고완료(s)",
           "첫 완료 ~ 마지막 완료", "@"),
          ("주문", "  └ 평균 총처리시간 (완료−발생)", "초", lambda c: f'AVERAGE({ocol("총처리시간(s)", c)})',
           -1, "주문별_CASE* · 총처리시간(s)", "참고", T),
          ("주문", "주문 개수", "건", lambda c: f'COUNTA({ocol("주문ID", c)})', 0, "주문별_CASE* · 행 수", "", N)]
    for label, field, unit in (("이동거리", "이동거리(m)", "m"), ("이동시간", "이동시간(s)", "sec"),
                               ("작업시간", "작업시간(s)", "sec"), ("대기시간", "대기시간(s)", "sec"),
                               ("유휴시간", "유휴시간(s)", "sec")):
        note = {"대기시간": "다른 AMR·병목 때문에 멈춘 시간 (재출발 지연 포함)",
                "유휴시간": "배정된 주문 없이 서 있던 시간",
                "작업시간": "적재·RFID·적치·피킹·하역·WMS·작업자 개입 대기"}.get(label, "")
        for amr in ("AMR1", "AMR2", "AMR3"):
            R.append(("AMR", f"AMR별 {label} — {amr}", unit, (lambda c, a=amr, fl=field: amr_sum(fl, c, a)),
                      -1 if label in ("이동거리", "이동시간", "대기시간") else 0, f"AMR별 · {field}", note if amr == "AMR1" else "",
                      T if unit != "m" else N))
        R.append(("AMR", f"AMR별 {label} — 합계", unit, (lambda c, fl=field: amr_sum(fl, c)),
                  -1 if label in ("이동거리", "이동시간", "대기시간") else 0, f"AMR별 · {field}", "", T if unit != "m" else N))
    R += [("충돌", "충돌 예상 횟수", "회",
           lambda c: (f'{SYS("경로 충돌 횟수", c)}+{SYS("정면 충돌 횟수", c)}' if c == "CASE1"
                      else SYS("예약테이블 사전 회피(충돌 예상)", c)), 0,
           "시스템요약", "CASE1 = 실제 발생한 경로충돌+정면충돌 / CASE2 = 예약테이블이 사전에 감지·회피한 수 (정의가 달라 개선율 미산출)", N),
          ("충돌", "  └ 실제 발생한 경로 충돌", "회", lambda c: SYS("경로 충돌 횟수", c), -1, "시스템요약 · 경로 충돌 횟수", "", N),
          ("충돌", "  └ 실제 발생한 정면 충돌", "회", lambda c: SYS("정면 충돌 횟수", c), -1, "시스템요약 · 정면 충돌 횟수", "", N),
          ("충돌", "충돌 회피를 위한 정지 횟수", "회", lambda c: SYS("정지 횟수", c), -1, "시스템요약 · 정지 횟수",
           "다른 AMR 때문에 멈춘 횟수 (예약 대기 포함)", N),
          ("병목", "병목구간 대기시간", "sec", lambda c: SYS("병목 대기시간(s)", c), -1, "시스템요약 · 병목 대기시간(s)",
           "공용구역 밖에서 구역 진입을 기다린 시간 합계", T)]
    for z in ("RFID", "중앙교차로", "좁은통로", "출고진입부"):
        R.append(("병목", f"  └ {z}", "sec", (lambda c, z=z: SYS(f"병목[{z}] 대기시간(s)", c)), -1,
                  f"시스템요약 · 병목[{z}] 대기시간(s)", "", T))
    R += [("병목", "병목구간 대기 AMR 수 (최대 동시)", "대", lambda c: SYS("병목구간 최대 동시 대기 AMR 수", c), -1,
           "시스템요약 · 병목구간 최대 동시 대기 AMR 수", "같은 순간 병목구간 앞에서 동시에 기다린 AMR 최대 대수", N)]
    for z in ("RFID", "중앙교차로", "좁은통로", "출고진입부"):
        R.append(("병목", f"  └ {z}", "대", (lambda c, z=z: SYS(f"병목[{z}] 최대 동시 대기 AMR 수", c)), -1,
                  f"시스템요약 · 병목[{z}] 최대 동시 대기 AMR 수", "", N))
    R += [("병목", "병목구간 대기 AMR 수 (대기 경험)", "대", lambda c: SYS("병목구간 대기 경험 AMR 수", c), 0,
           "시스템요약 · 병목구간 대기 경험 AMR 수", "한 번이라도 병목 대기를 겪은 AMR 대수 (3대 중)", N),
          ("병목", "  └ 병목 발생 횟수 (참고)", "회", lambda c: SYS("병목 횟수", c), -1, "시스템요약 · 병목 횟수",
           "대기 시작~진입까지 1회", N)]
    for amr in ("AMR1", "AMR2", "AMR3"):
        R.append(("배차", f"주문→AMR 배정 결과 — {amr}", "건",
                  (lambda c, a=amr: f'COUNTIF({ocol("배정 AMR", c)},"{a}")'), 0,
                  "주문별_CASE* · 배정 AMR", "주문별 배정 AMR 은 주문별 시트" if amr == "AMR1" else "", N))
    dkey = "배정 당시 AMR-입고장 거리(m)"
    R += [("배차", "배정 당시 AMR-작업 위치 거리 (평균)", "m", lambda c: f'AVERAGE({ocol(dkey, c)})', -1,
           "주문별_CASE* · 배정 당시 AMR-입고장 거리(m)", "작업 위치 = 입고장 (모든 주문의 시작점)", T),
          ("배차", "배정 당시 AMR-작업 위치 거리 (최대)", "m", lambda c: f'MAX({ocol(dkey, c)})', -1,
           "주문별_CASE* · 배정 당시 AMR-입고장 거리(m)", "", N),
          ("시스템", "전체 작업 완료시간", "sec", lambda c: SYS("완료시간(makespan,s)", c), -1, "시스템요약 · 완료시간(makespan,s)",
           "마지막 주문 출고 완료 시각", T),
          ("시스템", "시간당 처리 주문량", "건/h", lambda c: SYS("시간당 처리량(건/h)", c), +1, "시스템요약 · 시간당 처리량(건/h)", "", T),
          ("RFID", "인식 성공 횟수", "회", lambda c: SYS("RFID 인식 성공", c), 0, "시스템요약 · RFID 인식 성공",
           "RFID 결과는 주문 생성 시 고정 → 두 CASE 동일", N),
          ("RFID", "인식 실패 횟수 (감지)", "회", lambda c: SYS("RFID 인식 실패(감지·작업자 처리)", c), 0,
           "시스템요약 · RFID 인식 실패(감지·작업자 처리)", "작업자가 수동 확인 (10초)", N),
          ("RFID", "오인식 횟수 (미감지)", "회", lambda c: SYS("RFID 오인식(미감지)", c), 0, "시스템요약 · RFID 오인식(미감지)",
           "WMS 기록 불일치 → 재고 정확도에 반영", N)]
    first_int = len(R) + 1
    R += [("작업자", "작업자 개입 — 장시간 정체", "회", lambda c: SYS("장시간 정체 개입 횟수", c), -1, "시스템요약 · 장시간 정체 개입 횟수",
           "충돌·병목으로 연속 5초 이상 정지", N),
          ("작업자", "작업자 개입 — 수동 배차 확인", "회", lambda c: SYS("수동 배차 확인 개입 횟수", c), -1,
           "시스템요약 · 수동 배차 확인 개입 횟수", "배정 비용 ≥ 최소 비용 × 1.2", N),
          ("작업자", "작업자 개입 — RFID 오류", "회", lambda c: SYS("RFID 오류 작업자 개입 횟수", c), 0,
           "시스템요약 · RFID 오류 작업자 개입 횟수", "", N),
          ("작업자", "작업자 개입 횟수 — 합계", "회", None, -1, "위 3행 합계", "", N),
          ("작업자", "작업자 투입시간", "min", lambda c: f'{SYS("작업자 투입시간(s)", c)}/60', 0, "시스템요약 · 작업자 투입시간(s) ÷ 60",
           "태그 부착 + RFID 오류 처리 + 수동 배차 확인", T)]

    # ---- 결과정리 시트 작성
    ws = main
    ws["A1"] = "AMR 창고 시뮬레이션 — 필요 데이터 결과 정리 (Before CASE1 vs After CASE2)"
    ws["A1"].font = f(True, size=14)
    ws["A2"] = ("CASE1: 단순 순차 배정 + 독립 A*  ·  CASE2: OR-Tools VRP 배정 + Cooperative A* + 예약테이블  ·  "
                "AMR 3대, 주문 1,000건, seed 42 (두 CASE 동일 주문)")
    ws["A2"].font = f(color="555555", size=9)
    hdr = ["구분", "필요한 데이터", "단위", "CASE1 (Before)", "CASE2 (After)", "개선율", "출처 (시트 · 열)", "비고"]
    HR = 4
    for i, h in enumerate(hdr, 1):
        c = ws.cell(HR, i, h)
        c.font, c.fill, c.border = f(True, "FFFFFF"), HDR_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[HR].height = 24

    row = HR + 1
    group_start, prev = row, None
    int_rows = []
    for idx, (grp, item, unit, fn, direction, src, note, fmt) in enumerate(R, start=1):
        if grp != prev and prev is not None:
            if row - 1 > group_start:
                ws.merge_cells(start_row=group_start, start_column=1, end_row=row - 1, end_column=1)
            group_start = row
        if grp != prev:
            ws.cell(row, 1, grp)
        prev = grp
        ws.cell(row, 2, item)
        ws.cell(row, 3, unit)
        if fn is None:   # 작업자 개입 합계
            for col in (4, 5):
                L = get_column_letter(col)
                ws.cell(row, col, f"=SUM({L}{int_rows[0]}:{L}{int_rows[-1]})")
        else:
            ws.cell(row, 4, "=" + fn("CASE1"))
            ws.cell(row, 5, "=" + fn("CASE2"))
        if grp == "작업자" and fn is not None and "투입시간" not in item:
            int_rows.append(row)
        if direction == -1:
            ws.cell(row, 6, f'=IF(D{row}=0,"-",(D{row}-E{row})/D{row})')
        elif direction == 1:
            ws.cell(row, 6, f'=IF(D{row}=0,"-",(E{row}-D{row})/D{row})')
        else:
            ws.cell(row, 6, "-")
        ws.cell(row, 7, src)
        ws.cell(row, 8, note)
        sub = item.startswith("  └")
        for col in range(1, 9):
            c = ws.cell(row, col)
            c.border = BOX
            c.font = f(bold=(col == 1), size=10 if col != 7 else 9, color="000000" if col != 7 else "666666")
            c.alignment = Alignment(vertical="center", horizontal="center" if col in (1, 3, 6) else
                                    ("right" if col in (4, 5) else "left"), wrap_text=(col == 8))
            if col == 1:
                c.fill = GRP_FILL
            elif sub:
                c.fill = SUB_FILL
        for col in (4, 5):
            ws.cell(row, col).number_format = fmt if fmt != "@" else "General"
        ws.cell(row, 6).number_format = "+0.0%;-0.0%;0.0%"
        row += 1
    if row - 1 > group_start:
        ws.merge_cells(start_row=group_start, start_column=1, end_row=row - 1, end_column=1)

    notes = [
        "읽는 법",
        "· 개선율: 낮을수록 좋은 지표 (CASE1−CASE2)/CASE1, 시간당 처리량만 (CASE2−CASE1)/CASE1. 양수 = CASE2 가 좋음. '-' = 비교 대상 아님.",
        "· 모든 값은 시뮬레이션 실행값이며, 이 시트의 숫자는 시스템요약·AMR별·주문별 시트를 참조하는 수식입니다.",
        "· 병목구간 = RFID 구역·중앙교차로·좁은통로·출고진입부. 판정 기준: docs/new_scenario.md 4장.",
        "· 사람개입 판정 기준(5초·20%)과 작업자 개입 시간(10초)은 가정값 — amr_sim/params.py 에서 교체.",
    ]
    r0 = row + 1
    for i, t in enumerate(notes):
        c = ws.cell(r0 + i, 1, t)
        c.font = f(bold=(i == 0), size=9, color="333333")
    for col, w in zip("ABCDEFGH", (9, 38, 7, 17, 17, 10, 44, 52)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(HR + 1, 4)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", default="results")
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    out = a.out or os.path.join(a.results, "결과정리.xlsx")
    print("저장:", build(a.results, out))


if __name__ == "__main__":
    main()
